"""The user's OOXML answer sheet stays intact; working copy is conspicuously incomplete."""
from pathlib import Path

from docx import Document

from scripts.create_official_working_copy import TEMPLATE, make_working_copy


def test_university_sheet_is_word_and_all_subsections_remain(tmp_path):
    assert TEMPLATE.read_bytes()[:2] == b"PK"
    target = make_working_copy(TEMPLATE, tmp_path / "incomplete.docx")
    original, working = Document(TEMPLATE), Document(target)
    assert len(original.tables) == len(working.tables) == 38
    assert [p.text for p in original.paragraphs if p.style.name.startswith("Heading")] == [
        p.text for p in working.paragraphs if p.style.name.startswith("Heading")]
    assert "WORKING COPY — INCOMPLETE" in working.paragraphs[4].text
    assert working.tables[0].cell(0, 1).text == ""  # no fabricated name
    assert working.tables[1].cell(0, 0).text.startswith("[ Screenshot")  # not forged
    assert "MultilayerPerceptronClassifier" in working.tables[13].cell(3, 0).text
    assert working.tables[24].cell(1, 2).text == ""  # no invented stability delta
    assert working.tables[30].cell(1, 1).text == ""  # no invented Tableau URL


def test_source_only_zip_includes_real_sheet_but_no_student_results(tmp_path):
    from zipfile import ZipFile
    from scripts.package_project import build
    destination = tmp_path / "source-only.zip"
    report = build(destination)
    assert report["controlled_files"] >= 40 and report["bytes"] > 100_000
    with ZipFile(destination) as package:
        files = package.namelist()
        assert package.testzip() is None
        assert any(x.endswith("/report/7006SCN_Official_Answer_Sheet_TEMPLATE.docx") for x in files)
        assert any(x.endswith("/report/TR04_Official_Answer_Sheet_WORKING_INCOMPLETE.docx") for x in files)
        assert not any("/config/config.json" in x or "/results/task1.json" in x or "/.venv/" in x
                       for x in files)
