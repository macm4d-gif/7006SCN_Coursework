"""Assemble a Word DRAFT from real result JSON + the student's OWN written analysis.

This script intentionally refuses to produce an assessable-looking report until:
  * Tasks 1-4 have observed results, five real figures and live Tableau link;
  * private module-organisation git origin and cited commits exist;
  * every subsection contains the student's completed text (no placeholders);
  * a truthful AI declaration and references are supplied.
It is NOT the University's original answer-sheet template: review and transfer
its content to the provided sheet if the module requires that exact layout.
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.opc.constants import RELATIONSHIP_TYPE

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from coursework.settings import load_config, project_path, read_json, require_verified_allocation
from coursework.task4 import verify_public_link

SECTION_IDS = ("1.1", "1.2", "1.3", "EP1", "2.1", "2.2", "2.3", "EP2",
               "3.1", "3.2", "3.3", "4.1", "4.2")
SECTION_NAMES = {
    "1.1": "Problem, dataset and why this pairing", "1.2": "Five Vs and compliance",
    "1.3": "Data preparation and pipeline", "EP1": "Evidence pack 1",
    "2.1": "Algorithm selection", "2.2": "Tuning and compute", "2.3": "Metrics comparison",
    "EP2": "Evidence pack 2", "3.1": "Spark UI and optimisation",
    "3.2": "Stability on perturbed held-out test data", "3.3": "Explainability and fairness",
    "4.1": "Four dashboard contact sheet", "4.2": "Insights and critical reflection",
}
FIGURES = {"EP1": "task1_evidence.png", "EP2": "task2_evidence.png",
           "3.1": "task3_spark_ui.png", "3.3": "task3_lime.png",
           "4.1": "task4_dashboard_contact_sheet.png"}


def git_provenance(repository_url: str, sections: dict) -> None:
    expected = re.fullmatch(r"https://github\.com/7006SCN2627SEPNOV/7006SCN_[A-Za-z0-9]+_[0-9]+", repository_url)
    if not expected:
        raise ValueError("Report links must point to YOUR private 7006SCN_<Initials>_<SID> repo inside the module organisation")
    try:
        origin = subprocess.check_output(["git", "-C", str(ROOT), "remote", "get-url", "origin"],
                                         text=True, stderr=subprocess.DEVNULL).strip()
        all_hashes = subprocess.check_output(["git", "-C", str(ROOT), "log", "--format=%H"],
                                             text=True, stderr=subprocess.DEVNULL).splitlines()
    except (OSError, subprocess.CalledProcessError) as exc:
        raise ValueError("Cannot verify real Git history. Create/clone the required private module-org repository first.") from exc
    if not re.search(r"github\.com[:/]7006SCN2627SEPNOV/" + re.escape(repository_url.rsplit("/", 1)[-1]) + r"(?:\.git)?$", origin):
        raise ValueError(f"Git origin and report URL disagree or repository is outside module organisation: {origin}")
    for section in SECTION_IDS:
        commit = sections[section]["commit"]
        if not re.fullmatch(r"[0-9a-f]{7,40}", commit) or sum(h.startswith(commit) for h in all_hashes) != 1:
            raise ValueError(f"Section {section} must cite ONE real 7+ character short hash in THIS git log: {commit!r}")


def checked_author_content(path: Path, cfg: dict) -> dict:
    if not path.is_file():
        raise FileNotFoundError(f"Copy report/author_content.example.json to {path} and write YOUR analysis")
    content = json.loads(path.read_text(encoding="utf-8"))
    ai = content.get("ai_use_declaration", "")
    if len(ai.split()) < 25 or any(flag in ai.upper() for flag in ("WRITE_", "REPLACE", "TODO")):
        raise ValueError("Provide a meaningful truthful AI-use declaration; do not reuse the placeholder")
    sections = content.get("sections", {})
    for key in SECTION_IDS:
        section = sections.get(key, {})
        text = section.get("text", "")
        if len(text.split()) < 5 or any(flag in text.upper() for flag in ("WRITE_", "REPLACE", "TODO")):
            raise ValueError(f"Report {key}: use your own 5+ word evidence-based analysis, not a placeholder")
    references = content.get("references", [])
    if not references or any("ADD_" in item for item in references):
        raise ValueError("Add and check actual source references before generating the report")
    git_provenance(content.get("repository_url", ""), sections)
    return content


def hyperlink(paragraph, title: str, url: str) -> None:
    part = paragraph.part
    rid = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1264A3")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = title
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def table(doc, headers: list[str], rows: list[list[str]], counted: list[str]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for c, label in zip(t.rows[0].cells, headers):
        c.text = str(label)
    counted.extend(headers)
    for values in rows:
        for cell, value in zip(t.add_row().cells, values):
            cell.text = str(value)
            counted.append(str(value))
    doc.add_paragraph()


def safe_pic(doc, result_dir: Path, section: str) -> None:
    name = FIGURES.get(section)
    if name:
        path = result_dir / name
        if not path.is_file():
            raise FileNotFoundError(f"Required REAL evidence image missing: {path}")
        doc.add_picture(str(path), width=Inches(6.2))


def build(config_path: str, analysis_path: str, output_path: str) -> Path:
    cfg = load_config(config_path)
    require_verified_allocation(cfg)
    outdir = project_path(cfg, "results_dir")
    one, two, three, four = [read_json(outdir / f"task{i}.json") for i in (1, 2, 3, 4)]
    if any(result.get("status") != "observed" for result in (one, two, three, four)):
        raise ValueError("Need four observed result files from YOUR complete runs/published Tableau; placeholders are not acceptable")
    if one["pool_reference"] != cfg["allocation"]["pool_reference"] or len(two.get("models", [])) != 4:
        raise ValueError("Pool ID or four-model provenance mismatch")
    if not {"linear", "tree_ensemble", "neural"} <= {m.get("family") for m in two["models"]}:
        raise ValueError("Need four real results including the official answer sheet's neural third model family")
    if (three.get("test_perturbation_protocol", {}).get("labels_changed") != 0
            or len(three.get("stability", {})) != 4):
        raise ValueError("Need genuine four-model ΔF1/ΔROC-AUC on SAME perturbed TEST cohort")
    if not four.get("link_check", {}).get("valid") or len(four.get("dashboards", [])) != 4:
        raise ValueError("Four published dashboards and a verified live Tableau link are required")
    public_url = four["tableau_public_link"]
    if not verify_public_link(public_url)["valid"]:
        raise ValueError("The Tableau URL is no longer publicly accessible")
    content = checked_author_content(Path(analysis_path), cfg)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(9.5)
    counted: list[str] = []
    doc.add_heading("Report — 7006SCN Machine Learning and Big Data", 0)
    top_count = doc.add_paragraph("Word count (estimate incl. body tables/captions): calculating")
    doc.add_paragraph(f"Student: {cfg['student']['name']} | SID: {cfg['student']['sid']} | Email: {cfg['student']['email']}")
    doc.add_paragraph(f"Pool: {one['pool_reference']} | Source: {one['source_page']} | Terms: {one['licence_or_terms_url']}")
    doc.add_heading("AI Use Declaration", level=1)
    doc.add_paragraph(content["ai_use_declaration"])
    doc.add_paragraph("Repository: " + content["repository_url"])
    repo = content["repository_url"].rstrip("/")
    branch = content.get("branch", "main")
    if not re.fullmatch(r"[A-Za-z0-9_.-]+", branch):
        raise ValueError("Branch name contains unexpected characters")
    for n in (1, 2, 3, 4):
        hyperlink(doc.add_paragraph(), f"Task{n} Notebook (GitHub)", f"{repo}/blob/{branch}/notebooks/Task{n}.ipynb")
    hyperlink(doc.add_paragraph(), "Tableau Workbook (Public Link)", public_url)
    for key in SECTION_IDS:
        heading = f"{key} {SECTION_NAMES[key]}"
        doc.add_heading(heading, level=1 if key.startswith(("1", "2", "3", "4")) else 2)
        counted.append(heading)
        user_entry = content["sections"][key]
        for paragraph in user_entry["text"].split("\n\n"):
            doc.add_paragraph(paragraph)
            counted.append(paragraph)
        commit_note = f"Evidence commit: {user_entry['commit']}"
        doc.add_paragraph(commit_note, style="Caption")
        counted.append(commit_note)
        if key == "1.2":
            month_counts = list(one["raw_rows_by_source_month"].values())
            q = one["quality_counts_overlapping"]
            stats = [
                ["Volume", f"{one['raw_row_count']:,} rows; {one['file_size_gib']:.3f} GiB"],
                ["Velocity", f"{min(month_counts):,}–{max(month_counts):,} trips/month; historical files"],
                ["Variety", f"{min(m['original_column_count'] for m in one['raw_schema_by_month'])}–{max(m['original_column_count'] for m in one['raw_schema_by_month'])} raw fields/month"],
                ["Veracity", f"{q['invalid_fare']:,} invalid fares; {q['not_credit_card']:,} non-card/unobserved-tip records"],
                ["Value", "Decision only for trips with recorded electronic-payment tips"],
            ]
            table(doc, ["5 V", "Observed source evidence / conditional scope"], stats, counted)
        if key == "1.3":
            stages = one["preprocessing_stages"]
            table(doc, ["Stage order", "Actual Spark stage"], [[str(i), s] for i, s in enumerate(stages, 1)], counted)
        if key == "2.1":
            table(doc, ["Algorithm", "Family", "Grid candidates"],
                  [[m["name"], m["family"], str(m["param_grid"])] for m in two["models"]], counted)
        if key == "2.2":
            table(doc, ["Model", "Best params", "Folds", "CV / full-fit s"],
                  [[m["name"], str(m["best_params"]), str(m["cv_folds"]),
                    f"{m['cv_seconds']:.1f} / {m['final_fit_seconds']:.1f}"] for m in two["models"]], counted)
        if key == "2.3":
            table(doc, ["Model", "Acc/Bal/Spec", "ROC/PR AUC", "P/R/F1", "TN/FP/FN/TP"], [
                [m["name"],
                 f"{m['metrics']['accuracy']:.3f}/{m['metrics']['balanced_accuracy']:.3f}/{m['metrics']['specificity']:.3f}",
                 f"{m['metrics']['auc_roc']:.3f}/{m['metrics']['auc_pr']:.3f}",
                 f"{m['metrics']['positive_precision']:.3f}/{m['metrics']['positive_recall']:.3f}/{m['metrics']['positive_f1']:.3f}",
                 "/".join(str(m["metrics"]["confusion"][k]) for k in ("tn", "fp", "fn", "tp"))]
                for m in two["models"]], counted)
            baseline_note = (f"Nov–Dec high-tip prevalence / no-skill PR baseline: "
                             f"{two['models'][0]['metrics']['base_rate']:.3f}; "
                             f"CV-selected candidate: {two['recommended_model']}.")
            doc.add_paragraph(baseline_note)
            counted.append(baseline_note)
        if key == "3.1":
            optim = three["optimisations"]
            rows = []
            shuffles = [trial for trial in optim if trial["change"] == "shuffle_partition_count"]
            if len(shuffles) >= 2:
                rows.append([f"Shuffle {shuffles[0]['partitions']}→{shuffles[1]['partitions']}",
                             f"{shuffles[0]['median_seconds']:.2f}", f"{shuffles[1]['median_seconds']:.2f}",
                             "same groupBy; AQE may coalesce"])
            for row in optim:
                if row["change"] == "persist_MEMORY_AND_DISK":
                    rows.append(["Cache repeats", f"{row['median_before_seconds']:.2f}",
                                 f"{row['median_after_seconds']:.2f}", f"fill {row['cache_fill_seconds']:.2f}s"])
            table(doc, ["Change", "Before s", "After s", "Context"], rows, counted)
        if key == "3.2":
            test_protocol = three["test_perturbation_protocol"]
            table(doc, ["Model", "Test input perturbation", "Δ F1", "Δ AUC-ROC", "Rank"], [
                [name, test_protocol["perturbation"],
                 f"{entry['signed_deltas']['positive_f1']:+.4f}",
                 f"{entry['signed_deltas']['auc_roc']:+.4f}", str(entry["rank"])]
                for name, entry in sorted(three["stability"].items(),
                                          key=lambda pair: pair[1]["rank"])], counted)
        if key == "3.3":
            table(doc, ["LIME local feature", "Weight"], [
                [item["feature"], f"{item['local_weight']:+.4f}"]
                for item in three["explainability"]["feature_weights"][:5]], counted)
        safe_pic(doc, outdir, key)
    doc.add_heading("References", level=1)
    for item in content["references"]:
        doc.add_paragraph(item)
    body_words = len(re.findall(r"\b[\w]+(?:[-'.][\w]+)*\b", " ".join(counted)))
    top_count.text = f"Word count (estimated report body including tables and captions; excluding AI declaration/references): {body_words}"
    doc.add_paragraph(f"Word count (same estimate): {body_words}")
    if not 1080 <= body_words <= 1320:
        raise ValueError(f"Estimated body word count {body_words}; revise YOUR text/tables to reach 1080–1320, "
                         "then confirm actual count in the University's Word template. No DOCX was saved.")
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default="config/config.json")
    parser.add_argument("--analysis", default="report/author_content.json")
    parser.add_argument("--output", default="report/7006SCN_Report_DRAFT_review.docx")
    args = parser.parse_args()
    created = build(args.config, args.analysis, args.output)
    print(f"Built {created}. Review every figure, citation, link, AI declaration and the word count in Word.")
    print("Transfer to the University's supplied answer sheet if that specific template is required.")


if __name__ == "__main__":
    main()
