"""Prepare the USER-UPLOADED university answer sheet without inventing evidence.

The input was uploaded with a .pdf extension but is OOXML Word. The preserved
byte-for-byte copy lives at report/7006SCN_Official_Answer_Sheet_TEMPLATE.docx.
This working copy inserts only publicly verifiable/project-design information;
it deliberately leaves student identity, measurements, links, hashes, AI-use
sign-off, screenshots and checklist unchecked. NOT an assessed report.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "report" / "7006SCN_Official_Answer_Sheet_TEMPLATE.docx"
DESTINATION = ROOT / "report" / "TR04_Official_Answer_Sheet_WORKING_INCOMPLETE.docx"


def make_working_copy(template: Path = TEMPLATE, destination: Path = DESTINATION) -> Path:
    doc = Document(template)
    if len(doc.tables) != 38 or "1.1" not in "\n".join(p.text for p in doc.paragraphs):
        raise ValueError("The supplied university answer sheet has an unexpected structure; stop rather than removing sections")
    # Reuse an existing EMPTY title-page paragraph; add/remove NO subsections.
    if doc.paragraphs[4].text.strip():
        raise ValueError("Expected an empty title-page paragraph; do not overwrite assessment instructions")
    doc.paragraphs[4].text = (
        "WORKING COPY — INCOMPLETE; DO NOT SUBMIT. Finish all results, actual hyperlinks, "
        "eleven real section hashes, YOUR signed MS Forms screenshot, AI declaration and prompt log."
    )
    details = doc.tables[0]
    details.cell(2, 1).text = "TR-04 — user-stated; confirm against YOUR Aula allocation row"
    details.cell(3, 1).text = (
        "2019 NYC TLC Yellow Taxi Trip Records; "
        "https://www.nyc.gov/site/tlc/about/tlc-trip-record-data.page"
    )
    # Name the four IMPLEMENTED classifiers; their real grid winners/timings,
    # explanations and metric results must be supplied by observed runs.
    alg = doc.tables[13]
    for row, name, family in (
        (1, "LogisticRegression", "linear"),
        (2, "RandomForestClassifier", "tree ensemble"),
        (3, "MultilayerPerceptronClassifier", "neural"),
        (4, "GBTClassifier", "free choice — boosting"),
    ):
        alg.cell(row, 0).text = name
        alg.cell(row, 1).text = family
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    return destination


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", default=str(DESTINATION))
    args = parser.parse_args()
    path = make_working_copy(destination=Path(args.output))
    print(f"Created {path}; it is UNFINISHED and not submission-ready.")
