"""Fail closed on missing measured coursework evidence; NOT a grading predictor."""
from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path

import requests
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from coursework.task4 import verify_public_link

RESULTS = ROOT / "results"
parser = argparse.ArgumentParser()
parser.add_argument("--report", required=True,
                    help="Path to YOUR completed original-layout university answer-sheet .docx")
args = parser.parse_args()
errors, warnings = [], []


def error(message: str) -> None:
    errors.append(message)


def warn(message: str) -> None:
    warnings.append(message)


def result(task: int) -> dict:
    path = RESULTS / f"task{task}.json"
    if not path.is_file():
        error(f"Task{task}: real results JSON missing; run notebook on allocated dataset")
        return {}
    try:
        obj = json.loads(path.read_text(encoding="utf-8"))
        if obj.get("status") != "observed":
            error(f"Task{task}: result status {obj.get('status')!r} is not observed")
        return obj
    except (ValueError, OSError) as exc:
        error(f"Task{task}: cannot parse JSON: {exc}")
        return {}


for i in range(1, 5):
    path = ROOT / "notebooks" / f"Task{i}.ipynb"
    if not path.is_file():
        error(f"Task{i}: notebook missing")
    else:
        nb = json.loads(path.read_text(encoding="utf-8"))
        if not any(c.get("outputs") for c in nb.get("cells", []) if c.get("cell_type") == "code"):
            warn(f"Task{i}: notebook has no executed output; run and preserve evidence on real data")
if not (ROOT / "README.md").is_file():
    error("README missing")

one, two, three, four = [result(n) for n in (1, 2, 3, 4)]
if one:
    if one.get("raw_row_count", 0) < 10_000_000 or one.get("raw_column_count", 0) < 10 or one.get("file_size_bytes", 0) < 1024**3:
        error("Task1: observed dataset does not meet >=10M rows, >=10 columns, >=1 GiB")
    if (one.get("source_file_count") != 12 or len(one.get("preprocessing_stages", [])) < 3
            or one.get("preprocessing_fit_sample_rows", 0) < 100
            or not any(str(stage).endswith("Model") for stage in one.get("preprocessing_stages", []))):
        error("Task1: need all 12 official months and actually FITTED training-only pipeline stages for EP1")
    for key in ("pool_reference", "source_page", "licence_or_terms_url", "spark_configuration"):
        if not one.get(key):
            error(f"Task1: missing {key}")
    if not (RESULTS / "task1_evidence.png").is_file():
        error("Task1: actual evidence pack PNG missing")
if two:
    models = two.get("models", [])
    if (len(models) != 4 or len({m.get("name") for m in models}) != 4
            or not {"linear", "tree_ensemble", "neural"} <= {m.get("family") for m in models}):
        error("Task2: need four distinct algorithms including LINEAR, TREE ENSEMBLE and NEURAL families")
    for m in models:
        if m.get("cv_folds", 0) < 2 or not m.get("best_params") or m.get("cv_seconds", 0) <= 0:
            error(f"Task2: no real CV/folds/best grid/timing for {m.get('name')}")
        metrics = m.get("metrics", {})
        required = {"auc_pr", "auc_roc", "positive_f1", "accuracy", "positive_precision", "positive_recall", "confusion"}
        if required - metrics.keys():
            error(f"Task2: incomplete held-out metrics for {m.get('name')}")
    if len({m.get("metrics", {}).get("auc_pr") for m in models}) != len(models):
        warn("Task2: identical PR-AUC across models; verify model/prediction provenance")
    if not (RESULTS / "task2_evidence.png").is_file():
        error("Task2: actual composite evidence PNG missing")
if three:
    if not three.get("optimisations") or len(three.get("stability", {})) != 4:
        error("Task3: require measured optimisation and stability of all four models")
    protocol = three.get("test_perturbation_protocol", {})
    if (protocol.get("rows", 0) <= 0 or protocol.get("labels_changed") != 0
            or protocol.get("hour_shifted", 0) + protocol.get("borough_hidden", 0) <= 0):
        error("Task3: missing genuine shared Nov-Dec TEST-data feature perturbation")
    for model, measured in three.get("stability", {}).items():
        delta = measured.get("signed_deltas", {})
        if ("positive_f1" not in delta or "auc_roc" not in delta
                or measured.get("rank") not in range(1, 5)):
            error(f"Task3: {model} needs ΔF1, ΔROC-AUC and stability rank on perturbed TEST data")
    lime = three.get("explainability", {})
    if lime.get("method", "").split(" ")[0] not in {"LIME", "SHAP"} or len(lime.get("feature_weights", [])) < 5:
        error("Task3: require actual Spark-model LIME/SHAP with >=5 features")
    if not (RESULTS / "task3_spark_ui.png").is_file():
        error("Task3: genuine Spark UI screenshot missing")
    if not (RESULTS / "task3_lime.png").is_file():
        error("Task3: genuine explanation figure missing")
if four:
    if not four.get("link_check", {}).get("valid") or len(four.get("dashboards", [])) != 4:
        error("Task4: live Tableau Public link/four-dashboard confirmation missing")
    elif not verify_public_link(four.get("tableau_public_link", "")).get("valid"):
        error("Task4: saved Tableau URL is no longer publicly accessible")
    if not (RESULTS / "task4_dashboard_contact_sheet.png").is_file():
        error("Task4: genuine four-dashboard contact sheet missing")

try:
    remote = subprocess.check_output(["git", "-C", str(ROOT), "remote", "get-url", "origin"], text=True, stderr=subprocess.DEVNULL).strip()
except (subprocess.CalledProcessError, FileNotFoundError):
    remote = ""
if remote:
    if not re.search(r"github\.com[:/]7006SCN2627SEPNOV/7006SCN_[^/]+(?:\.git)?$", remote):
        error(f"Origin outside module org or wrong repo name: {remote} (20% cap)")
    else:
        repository_name = re.search(r"/([^/]+?)(?:\.git)?$", remote).group(1)
        api = f"https://api.github.com/repos/7006SCN2627SEPNOV/{repository_name}"
        headers = {"Accept": "application/vnd.github+json"}
        token = os.getenv("GH_TOKEN")  # optional in ENV only; never store tokens in git.
        if token:
            headers["Authorization"] = f"Bearer {token}"
        try:
            response = requests.get(api, headers=headers, timeout=12)
            if response.ok and response.json().get("private") is False:
                error("GitHub repository is PUBLIC: brief deducts 10 marks. Set PRIVATE and check your module-team access.")
            elif response.ok and response.json().get("private") is True:
                print("GitHub API: private organisation repository confirmed")
            else:
                warn(f"GitHub API returned {response.status_code}; cannot verify private repo visibility without authenticated access")
        except requests.RequestException:
            warn("Could not query GitHub visibility; manually confirm private from the repository Settings page")
else:
    error("No git origin; create private repo INSIDE 7006SCN2627SEPNOV (not personal account)")
try:
    lines = subprocess.check_output(["git", "-C", str(ROOT), "log", "--format=%H|%as|%s"], text=True, stderr=subprocess.DEVNULL).splitlines()
except (subprocess.CalledProcessError, FileNotFoundError):
    lines = []
counts = Counter()
for line in lines:
    fields = line.split("|", 2)
    if len(fields) != 3:
        continue
    match = re.match(r"^\[Task([1-4])\] [a-z]+ .+ - .+$", fields[2])
    if match:
        counts[int(match.group(1))] += 1
for task in range(1, 5):
    if counts[task] < 3:
        error(f"Task{task}: only {counts[task]} correctly formatted actual commits (need >=3)")
report_path = Path(args.report)
if not report_path.is_file():
    error(f"Single Word report missing: {report_path} (use --report path/to/official_answer_sheet.docx)")
else:
    try:
        document = Document(report_path)
        text = "\n".join(p.text for p in document.paragraphs)
        if not re.search(r"AI[-\s]Use Declaration", text, flags=re.IGNORECASE):
            error("Report lacks the official AI-Use Declaration")
        required_answer_sheet_sections = (
            "Signed Checklist Screenshot", "Student Details", "Before You Submit",
            "Appendix — Prompt Log", "1.1", "1.2", "1.3", "2.1", "2.2", "2.3",
            "3.1", "3.2", "3.3", "4.1", "4.2",
        )
        for section in required_answer_sheet_sections:
            if section not in text:
                error(f"Use and COMPLETE the supplied university answer sheet; missing {section!r}")
        checklist_tables = [tbl for tbl in document.tables if any(
            "Screenshot of completed & signed MS Forms checklist" in cell.text
            for row in tbl.rows for cell in row.cells)]
        if checklist_tables:
            error("Report still contains the BLANK signed-checklist screenshot placeholder")
        if len(document.tables) < 38:
            error("Report is missing original answer-sheet tables; preserve every subsection/table")
        elif not (document.tables[1].cell(0, 0)._element.xpath(".//w:drawing")
                  or document.tables[1].cell(0, 0)._element.xpath(".//w:pict")):
            error("Report lacks an actual image in the signed MS Forms checklist cell")
        if len(document.tables) >= 38:
            details = document.tables[0]
            for index, label in ((0, "full name"), (1, "sid"), (2, "pool"), (3, "dataset")):
                value = details.cell(index, 1).text.strip()
                if not value or "Student-Name-Mail-Dataset" in value:
                    error(f"Official answer sheet: {label} field in Student Details is blank")
            ai_table = document.tables[2]
            if not any(any(c.text.strip() for c in row.cells) for row in ai_table.rows[1:]):
                error("AI-Use Declaration table is BLANK; declare Arena.ai use and actual prompts")
            if not document.tables[3].cell(0, 1).text.strip():
                error("AI-use confirmation has no student's typed name/signature")
            if not any(row.cells[0].text.strip() for row in document.tables[36].rows):
                error("Official answer sheet: Appendix Prompt Log is BLANK")
            count_text = details.cell(7, 2).text.strip()
        else:
            count_text = ""
        reported = re.search(r"^\s*([0-9]{3,4})\s*$", count_text)
        if reported is None:
            reported = re.search(r"Total (?:body )?word count\s*:\s*([0-9]{3,4})",
                                 text.split("Task 1 —")[0], flags=re.IGNORECASE)
        if reported:
            n = int(reported.group(1))
            if not 1080 <= n <= 1320:
                error(f"Declared body word count {n} is outside the safe 1080–1320 band")
        else:
            error("Enter 1080–1320 total body words in Student Details at the START of the answer sheet")
        targets = [rel.target_ref for rel in document.part.rels.values()
                   if rel.reltype == RELATIONSHIP_TYPE.HYPERLINK]
        for n in (1, 2, 3, 4):
            if not any(f"/notebooks/Task{n}.ipynb" in href and
                       "github.com/7006SCN2627SEPNOV/" in href for href in targets):
                error(f"Report: accessible module-organisation link for Task{n} notebook missing")
        if four and four.get("tableau_public_link") not in targets:
            error("Report: active Tableau Public workbook hyperlink missing or differs from task4.json")
        real_hashes = [field.split("|", 1)[0] for field in lines]
        cited = re.findall(r"(?:Evidence commit|Commit hash):\s*([0-9a-f]{7,40})", text,
                           flags=re.IGNORECASE)
        if len(cited) < 11 or any(not any(h.startswith(short) for h in real_hashes) for short in cited):
            error("Report must cite at least 11 genuine local commit hashes in its sub-sections")
    except Exception as exc:
        error(f"Could not validate Word report: {type(exc).__name__}: {exc}")
warn("Manually verify a measured-outcome commit per Task, correct teaching-week timestamps, marker access to all private notebook links and data allocation")
warn("Review actual word count in Word, usage terms and the truthful AI declaration; automatic checks do NOT establish authorship or a mark")
print("CHECK SUMMARY")
for problem in errors:
    print("FAIL:", problem)
for message in warnings:
    print("CHECK:", message)
print(f"{len(errors)} blocking checks; {len(warnings)} manual checks. A zero check count is NOT a mark guarantee.")
sys.exit(1 if errors else 0)
